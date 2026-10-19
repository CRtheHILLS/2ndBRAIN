import csv, hashlib, re, datetime as dt
from io import BytesIO
from pathlib import Path
import pymupdf
from PIL import Image
from .config import get_settings
from . import catalog, llm, ocr, store  # noqa: F401  (llm kept: tests patch ingest.llm)

fitz = pymupdf

IMG = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
       ".heic": "image/heic", ".heif": "image/heif"}
HEIC_EXTS = {".heic", ".heif"}

def _pdf_text(data: bytes) -> tuple[str, bool]:
    """Return (text, ocr_incomplete). Pages that fail OCR get a placeholder comment."""
    out, incomplete = [], False
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, 1):
            t = page.get_text().strip()
            if not t:  # scanned page -> OCR
                try:
                    t = ocr.ocr_image(page.get_pixmap(dpi=150).tobytes("png"), "image/png")
                except ocr.OcrUnavailable:
                    incomplete = True
                    out.append(f"\n\n<!-- page {i}: ocr pending -->")
                    continue
            out.append(f"\n\n<!-- page {i} -->\n{t}")
    return "".join(out), incomplete

def _to_jpeg(data: bytes, ext: str) -> bytes:
    """Convert image bytes (incl. HEIC/HEIF) to JPEG bytes for OCR."""
    if ext in HEIC_EXTS:
        import pillow_heif
        pillow_heif.register_heif_opener()
    img = Image.open(BytesIO(data)).convert("RGB")
    buf = BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()

def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return data.decode("cp949")
        except UnicodeDecodeError:
            return data.decode("utf-8", errors="replace")

def _escape_cell(v) -> str:
    s = "" if v is None else str(v)
    s = s.replace("|", "\\|")
    for nl in ("\r\n", "\r", "\n"):
        s = s.replace(nl, "<br>")
    return s

def _md_table(rows: list) -> str:
    rows = [list(r) for r in rows]
    if not rows:
        return ""
    header, *body = rows
    header = [_escape_cell(c) for c in header]
    lines = ["| " + " | ".join(header) + " |",
              "| " + " | ".join(["---"] * len(header)) + " |"]
    for row in body:
        lines.append("| " + " | ".join(_escape_cell(c) for c in row) + " |")
    return "\n".join(lines)

def _xlsx_tables(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        parts = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            parts.append(f"## {ws.title}\n\n{_md_table(rows)}")
        return "\n\n".join(parts)
    finally:
        wb.close()

def _csv_table(data: bytes, stem: str) -> str:
    text = _decode_text(data)
    rows = list(csv.reader(text.splitlines()))
    return f"## {stem}\n\n{_md_table(rows)}"

_HEADING_RE = re.compile(r"^(Heading|제목)\s*(\d+)$")

def _iter_docx_blocks(document):
    """Yield paragraphs and tables in document order."""
    import docx.table
    import docx.text.paragraph
    from docx.oxml.ns import qn
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield docx.text.paragraph.Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield docx.table.Table(child, document)

def _docx_body(data: bytes) -> str:
    import docx
    document = docx.Document(BytesIO(data))
    parts = []
    for block in _iter_docx_blocks(document):
        if hasattr(block, "rows"):  # Table
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            table_md = _md_table(rows)
            if table_md:
                parts.append(table_md)
            continue
        text = block.text
        if not text.strip():
            continue
        style = (block.style.name if block.style else "") or ""
        m = _HEADING_RE.match(style)
        if m:
            n = min(int(m.group(2)), 6)
            parts.append("#" * n + " " + text)
        elif "List" in style or "목록" in style:
            parts.append("- " + text)
        else:
            parts.append(text)
    return "\n\n".join(parts)

def _sanitize_filename(filename: str) -> str:
    name = filename.replace("\\", "/").split("/")[-1]
    name = Path(name).name
    if not name or name in ("..", "."):
        raise ValueError(f"invalid filename: {filename!r}")
    return name


def _catalog_register(slug: str, filename: str, out: Path, sha: str,
                      hold_until: str | None = None) -> None:
    """Put the file on the knowledge-catalog learn queue (``hold_until`` = grace period). Never fail ingest."""
    try:
        from . import catalog
        catalog.register(slug, filename, out, sha, hold_until=hold_until)
    except Exception:
        pass


def ingest_file(slug: str, filename: str, data: bytes, extra_meta: dict | None = None,
                hold_until: str | None = None) -> Path:
    """Store the original, write the markdown. Never raises OcrUnavailable: the md is marked
    ``ocr: pending`` and queued for Clair to transcribe."""
    filename = _sanitize_filename(filename)
    sha = hashlib.sha256(data).hexdigest()
    ext = Path(filename).suffix.lower()
    out = store.raw_dir(slug) / (Path(filename).stem + ".md")
    prior: dict = {}
    if out.exists() and store.read_md(out)[0].get("sha256") == sha:
        old = store.read_md(out)[0]
        if old.get("ocr") != "pending":
            _catalog_register(slug, filename, out, sha, hold_until)
            return out
        eid = catalog.make_id(slug, filename)
        if any(r["id"] == eid for r in ocr.list_pending()):
            return out
        prior = old  # pending md lost its queue entry: retry OCR and re-queue
    (store.raw_dir(slug) / filename).write_bytes(data)  # keep original, before any llm call
    ocr_state, pending_mime = "none", None
    if ext in HEIC_EXTS or ext in IMG:
        kind = "image"
        jpeg = ext in HEIC_EXTS
        mime = "image/jpeg" if jpeg else IMG[ext]
        try:
            if (extra_meta or {}).get("via") == "telegram" and get_settings().ocr_mode == "clair":
                raise ocr.OcrUnavailable(["clair-first"])  # Clair reads Telegram photos herself
            body = ocr.ocr_image(_to_jpeg(data, ext) if jpeg else data, mime)
            ocr_state = "done"
        except ocr.OcrUnavailable:
            body, ocr_state, pending_mime = "", "pending", IMG.get(ext, mime)
    elif ext == ".pdf":
        kind = "pdf"
        body, incomplete = _pdf_text(data)
        ocr_state = "pending" if incomplete else "none"
        pending_mime = "application/pdf"
    elif ext == ".xlsx":
        kind, body = "table", _xlsx_tables(data)
    elif ext == ".csv":
        kind, body = "table", _csv_table(data, Path(filename).stem)
    elif ext == ".docx":
        kind, body = "doc", _docx_body(data)
    else:
        kind, body = "text", _decode_text(data)
    core = {"book": slug, "source": filename, "sha256": sha, "kind": kind, "ocr": ocr_state,
            "ingested_at": dt.datetime.now().isoformat(timespec="seconds")}
    meta = {**prior, **(extra_meta or {}), **core}  # core keys always win
    store.write_md(out, meta, body)
    if ocr_state == "pending":
        ocr.add_pending({"id": catalog.make_id(slug, filename), "book": slug, "file": filename,
                         "path": str(out), "image": str(store.raw_dir(slug) / filename),
                         "mime": pending_mime})
    else:
        _catalog_register(slug, filename, out, sha, hold_until)
    return out
