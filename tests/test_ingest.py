from io import BytesIO
import docx
import openpyxl
import pymupdf
from PIL import Image
from brain import ingest, store

def test_text_passthrough(data_dir):
    p = ingest.ingest_file("cosmos", "ch1.txt", "우주는 넓다".encode())
    meta, body = store.read_md(p)
    assert meta["kind"] == "text" and "우주는 넓다" in body

def _anthropic_key(monkeypatch):
    from brain.config import get_settings
    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    get_settings.cache_clear()

def test_image_uses_ocr(data_dir, monkeypatch):
    _anthropic_key(monkeypatch)
    monkeypatch.setattr(ingest.llm, "ocr_image", lambda b, m: "OCR 결과 ==밑줄==")
    p = ingest.ingest_file("cosmos", "IMG_1.jpg", b"\xff\xd8fake")
    assert "==밑줄==" in store.read_md(p)[1]

def test_idempotent(data_dir):
    a = ingest.ingest_file("cosmos", "a.txt", b"same")
    b = ingest.ingest_file("cosmos", "a.txt", b"same")
    assert a == b

def test_path_traversal_filename_is_sanitized(data_dir):
    p = ingest.ingest_file("cosmos", "../../evil.txt", b"pwned")
    assert p.name == "evil.md"
    assert p.parent == store.raw_dir("cosmos")
    assert (store.raw_dir("cosmos") / "evil.txt").exists()
    assert not (data_dir / "evil.txt").exists()

def test_backslash_path_traversal_filename_is_sanitized(data_dir):
    p = ingest.ingest_file("cosmos", "..\\..\\evil2.txt", b"pwned")
    assert p.parent == store.raw_dir("cosmos")
    assert (store.raw_dir("cosmos") / "evil2.txt").exists()

def test_dotdot_filename_rejected(data_dir):
    import pytest
    with pytest.raises(ValueError):
        ingest.ingest_file("cosmos", "..", b"x")

def test_empty_filename_rejected(data_dir):
    import pytest
    with pytest.raises(ValueError):
        ingest.ingest_file("cosmos", "", b"x")

def test_raw_bytes_written_before_llm_call(data_dir, monkeypatch):
    _anthropic_key(monkeypatch)
    def _boom(b, m):
        raise RuntimeError("llm exploded")
    monkeypatch.setattr(ingest.llm, "ocr_image", _boom)
    ingest.ingest_file("cosmos", "IMG_2.jpg", b"\xff\xd8fake")  # provider error -> pending, no raise
    assert (store.raw_dir("cosmos") / "IMG_2.jpg").exists()

def test_cp949_text_is_decoded(data_dir):
    data = "우주는 넓다".encode("cp949")
    p = ingest.ingest_file("cosmos", "cp949.txt", data)
    meta, body = store.read_md(p)
    assert meta["kind"] == "text" and "우주는 넓다" in body

def test_to_jpeg_converts_png_bytes_to_jpeg_magic_bytes(data_dir):
    buf = BytesIO()
    Image.new("RGB", (4, 4), color="red").save(buf, "PNG")
    out = ingest._to_jpeg(buf.getvalue(), ".png")
    assert out[:2] == b"\xff\xd8"

def test_pdf_ingest_extracts_text_and_ocrs_blank_page(data_dir, monkeypatch):
    _anthropic_key(monkeypatch)
    calls = []
    monkeypatch.setattr(ingest.llm, "ocr_image", lambda b, m: (calls.append(1), "OCR")[1])
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "Hello Test Page")
    doc.new_page()  # blank page -> OCR
    pdf_bytes = doc.tobytes()
    doc.close()
    p = ingest.ingest_file("cosmos", "book.pdf", pdf_bytes)
    meta, body = store.read_md(p)
    assert meta["kind"] == "pdf"
    assert "Hello Test Page" in body
    assert "OCR" in body
    assert len(calls) == 1

def test_xlsx_ingest_produces_markdown_tables(data_dir):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "용어"
    ws1.append(["용어", "풀이"])
    ws1.append(["a|b", "line1\nline2"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["x", "y"])
    ws2.append([1, 2])
    buf = BytesIO()
    wb.save(buf)
    p = ingest.ingest_file("cosmos", "glossary.xlsx", buf.getvalue())
    meta, body = store.read_md(p)
    assert meta["kind"] == "table"
    assert "## 용어" in body
    assert "## Sheet2" in body
    assert "a\\|b" in body
    assert "line1<br>line2" in body
    assert (store.raw_dir("cosmos") / "glossary.xlsx").exists()

def test_csv_ingest_produces_markdown_table_with_korean_utf8(data_dir):
    csv_bytes = "용어,풀이\n우주,넓다\n".encode("utf-8")
    p = ingest.ingest_file("cosmos", "terms.csv", csv_bytes)
    meta, body = store.read_md(p)
    assert meta["kind"] == "table"
    assert "## terms" in body
    assert "| 용어 | 풀이 |" in body
    assert "| 우주 | 넓다 |" in body

def test_docx_ingest_renders_heading_and_table(data_dir):
    document = docx.Document()
    document.add_heading("우주론 개요", level=1)
    document.add_paragraph("우주는 넓다.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "용어"
    table.cell(0, 1).text = "풀이"
    table.cell(1, 0).text = "우주"
    table.cell(1, 1).text = "넓다"
    buf = BytesIO()
    document.save(buf)
    p = ingest.ingest_file("cosmos", "notes.docx", buf.getvalue())
    meta, body = store.read_md(p)
    assert meta["kind"] == "doc"
    assert "# 우주론 개요" in body
    assert "| 용어 | 풀이 |" in body
    assert (store.raw_dir("cosmos") / "notes.docx").exists()


def test_hold_until_is_forwarded_to_catalog(data_dir):
    import datetime as dt
    from brain import catalog
    later = (dt.datetime.now() + dt.timedelta(minutes=30)).isoformat(timespec="seconds")
    ingest.ingest_file("b", "note.txt", "안녕".encode(), hold_until=later)
    assert catalog.pending() == []
    (item,) = catalog.pending(include_held=True)
    assert item["held"] is True and item["hold_until"] == later
